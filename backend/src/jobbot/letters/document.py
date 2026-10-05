"""Assemblage de la lettre : coordonnées, lieu et date, destinataire, objet, corps (docs/08 §3).

L'IA n'écrit que l'objet et le corps ; tout le reste vient d'ici, à partir des coordonnées
saisies dans les Réglages et de l'offre. Le même assemblage sert à l'affichage et au Word.
"""

import io
from dataclasses import dataclass, field
from datetime import date

from docx import Document as new_docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

MONTHS = {
    "fr": [
        "janvier", "février", "mars", "avril", "mai", "juin",
        "juillet", "août", "septembre", "octobre", "novembre", "décembre",
    ],
    "en": [
        "January", "February", "March", "April", "May", "June",
        "July", "August", "September", "October", "November", "December",
    ],
    "de": [
        "Januar", "Februar", "März", "April", "Mai", "Juni",
        "Juli", "August", "September", "Oktober", "November", "Dezember",
    ],
}  # fmt: skip

SALUTATION = {
    "fr": "Madame, Monsieur,",
    "en": "Dear Sir or Madam,",
    "de": "Sehr geehrte Damen und Herren",
}
CLOSING = {
    "fr": "Je vous prie d'agréer, Madame, Monsieur, mes salutations distinguées.",
    "en": "Yours faithfully,",
    "de": "Freundliche Grüsse",
}
SUBJECT_PREFIX = {"fr": "Objet : ", "en": "Subject: ", "de": ""}
ENCLOSURE = {"fr": "Annexe : curriculum vitae", "en": "Enclosure: CV", "de": "Beilage: Lebenslauf"}
ATTENTION = {"fr": "À l'attention de ", "en": "For the attention of ", "de": "z. H. "}


@dataclass(frozen=True)
class Identity:
    name: str | None = None
    street: str | None = None
    postcode: str | None = None
    city: str | None = None
    phone: str | None = None
    email: str | None = None

    @property
    def missing(self) -> list[str]:
        """Champs nécessaires à l'en-tête et pas encore saisis."""
        labels = {"name": "nom", "street": "rue", "postcode": "NPA", "city": "localité"}
        return [label for key, label in labels.items() if not getattr(self, key)]


@dataclass(frozen=True)
class Recipient:
    company: str | None
    address: str | None = None
    contact_name: str | None = None


@dataclass
class LetterDocument:
    language: str
    sender: list[str]
    place_date: str
    recipient: list[str]
    subject_line: str
    salutation: str
    paragraphs: list[str]
    closing: str
    signature: str
    enclosure: str
    missing_identity: list[str] = field(default_factory=list)


def format_date(day: date, language: str) -> str:
    month = MONTHS[language][day.month - 1]
    if language == "fr":
        return f"{'1er' if day.day == 1 else day.day} {month} {day.year}"
    if language == "de":
        return f"{day.day}. {month} {day.year}"
    return f"{day.day} {month} {day.year}"


def assemble(
    identity: Identity,
    recipient: Recipient,
    language: str,
    subject: str,
    paragraphs: list[str],
    day: date,
) -> LetterDocument:
    city_line = " ".join(x for x in (identity.postcode, identity.city) if x)
    sender = [
        x for x in (identity.name, identity.street, city_line, identity.phone, identity.email) if x
    ]
    when = format_date(day, language)
    if identity.city:
        place_date = (
            f"{identity.city}, le {when}" if language == "fr" else f"{identity.city}, {when}"
        )
    else:
        place_date = f"Le {when}" if language == "fr" else when
    lines = [recipient.company or ""]
    if recipient.contact_name:
        lines.append(ATTENTION[language] + recipient.contact_name)
    if recipient.address:
        lines += [
            part.strip()
            for part in recipient.address.replace(",", "\n").splitlines()
            if part.strip()
        ]
    return LetterDocument(
        language=language,
        sender=sender,
        place_date=place_date,
        recipient=[line for line in lines if line],
        subject_line=SUBJECT_PREFIX[language] + subject,
        salutation=SALUTATION[language],
        paragraphs=paragraphs,
        closing=CLOSING[language],
        signature=identity.name or "",
        enclosure=ENCLOSURE[language],
        missing_identity=identity.missing,
    )


def to_docx(letter: LetterDocument) -> bytes:
    """Lettre A4 au format Word, mise en page suisse : destinataire à droite."""
    doc = new_docx()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2.5))
    section.top_margin = section.bottom_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(0)

    def block(lines: list[str], *, indent: bool = False, after: int = 18) -> None:
        if not lines:
            return
        p = doc.add_paragraph("\n".join(lines))
        if indent:
            p.paragraph_format.left_indent = Cm(9)
        p.paragraph_format.space_after = Pt(after)

    block(letter.sender)
    block(letter.recipient, indent=True)
    block([letter.place_date], indent=True, after=24)
    subject = doc.add_paragraph()
    subject.add_run(letter.subject_line).bold = True
    subject.paragraph_format.space_after = Pt(18)
    block([letter.salutation], after=12)
    for text in letter.paragraphs:
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(10)
    block([letter.closing], after=36)
    block([letter.signature], after=30)
    block([letter.enclosure], after=0)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
