"""Mise en page du CV adapté : une page A4 sobre (docs/08 §4).

Coordonnées, titre, résumé, puis expériences, compétences, formation et langues. Le texte des
blocs est repris tel quel ; seuls les mots-clés de l'offre qu'il contient sont mis en gras.
"""

import io
import re
from dataclasses import dataclass, field

from docx import Document as new_docx
from docx.shared import Cm, Pt, RGBColor

from jobbot.letters.document import Identity
from jobbot.llm.cv import is_language_block
from jobbot.llm.scoring import ProfileChunkData

HEADINGS = {
    "fr": {
        "summary": "Profil",
        "experience": "Expériences",
        "competence": "Compétences",
        "formation": "Formation",
        "langues": "Langues",
    },
    "en": {
        "summary": "Profile",
        "experience": "Experience",
        "competence": "Skills",
        "formation": "Education",
        "langues": "Languages",
    },
    "de": {
        "summary": "Profil",
        "experience": "Berufserfahrung",
        "competence": "Kompetenzen",
        "formation": "Ausbildung",
        "langues": "Sprachen",
    },
}
ORDER = ("experience", "competence", "formation", "langues")


@dataclass
class Segment:
    text: str
    strong: bool = False


@dataclass
class CvItem:
    chunk_id: int
    title: str | None
    content: list[Segment]


@dataclass
class CvSection:
    key: str
    heading: str
    items: list[CvItem]


@dataclass
class CvDocument:
    language: str
    name: str
    headline: str
    contacts: list[str]
    summary_heading: str
    summary: str
    sections: list[CvSection]
    missing_identity: list[str] = field(default_factory=list)


def short_url(url: str | None) -> str | None:
    """« https://www.linkedin.com/in/x/ » devient « linkedin.com/in/x » sur le CV."""
    if not url:
        return None
    return re.sub(r"^https?://(www\.)?", "", url.strip()).rstrip("/") or None


def section_of(chunk: ProfileChunkData) -> str:
    return "langues" if is_language_block(chunk) else chunk.kind


def highlight(text: str, keywords: list[str]) -> list[Segment]:
    """Découpe le texte en segments, les mots-clés (sans tenir compte de la casse) en gras."""
    words = sorted({k for k in keywords if k.strip()}, key=len, reverse=True)
    if not words:
        return [Segment(text)]
    pattern = re.compile(
        r"(?<!\w)(" + "|".join(re.escape(w) for w in words) + r")(?!\w)", re.IGNORECASE
    )
    segments: list[Segment] = []
    last = 0
    for match in pattern.finditer(text):
        if match.start() > last:
            segments.append(Segment(text[last : match.start()]))
        segments.append(Segment(match.group(0), strong=True))
        last = match.end()
    if last < len(text):
        segments.append(Segment(text[last:]))
    return segments


def assemble_cv(
    identity: Identity,
    chunks: dict[int, ProfileChunkData],
    chunk_ids: list[int],
    *,
    language: str,
    headline: str,
    summary: str,
    keywords: list[str],
) -> CvDocument:
    headings = HEADINGS[language]
    grouped: dict[str, list[CvItem]] = {key: [] for key in ORDER}
    for chunk_id in chunk_ids:
        chunk = chunks.get(chunk_id)
        if chunk is None:  # bloc supprimé ou désactivé depuis
            continue
        key = section_of(chunk)
        if key not in grouped:
            continue
        title = None if key == "langues" else chunk.title
        grouped[key].append(CvItem(chunk.id, title, highlight(chunk.content, keywords)))
    city = " ".join(x for x in (identity.postcode, identity.city) if x)
    address = ", ".join(x for x in (identity.street, city) if x)
    return CvDocument(
        language=language,
        name=identity.name or "",
        headline=headline,
        contacts=[
            x for x in (address, identity.phone, identity.email, short_url(identity.linkedin)) if x
        ],
        summary_heading=headings["summary"],
        summary=summary,
        sections=[CvSection(key, headings[key], grouped[key]) for key in ORDER if grouped[key]],
        missing_identity=identity.missing,
    )


ACCENT = RGBColor(0x3A, 0x33, 0xC9)
MUTED = RGBColor(0x55, 0x55, 0x66)


def cv_to_docx(cv: CvDocument) -> bytes:
    doc = new_docx()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)

    name = doc.add_paragraph()
    run = name.add_run(cv.name)
    run.bold, run.font.size = True, Pt(20)
    if cv.headline:
        headline = doc.add_paragraph()
        run = headline.add_run(cv.headline)
        run.font.size, run.font.color.rgb = Pt(12), ACCENT
    contacts = doc.add_paragraph()
    run = contacts.add_run(" · ".join(cv.contacts))
    run.font.size, run.font.color.rgb = Pt(9), MUTED
    contacts.paragraph_format.space_after = Pt(8)

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text.upper())
        run.bold, run.font.size, run.font.color.rgb = True, Pt(10.5), ACCENT

    if cv.summary:
        heading(cv.summary_heading)
        doc.add_paragraph(cv.summary)
    for block in cv.sections:
        heading(block.heading)
        for item in block.items:
            if item.title:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(3)
                p.add_run(item.title).bold = True
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            for segment in item.content:
                p.add_run(segment.text).bold = segment.strong
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
