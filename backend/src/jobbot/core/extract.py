"""Reconnaissance du type d'un document et extraction de son texte (docs/04 §5).

Types acceptés : PDF, DOCX, TXT, MD, 10 Mo au plus. Le type est vérifié sur le contenu
(signature du fichier), pas seulement sur l'extension. Pas de reconnaissance de texte pour
les PDF scannés : ils sont signalés « texte non lisible ».
"""

import io
import zipfile
from dataclasses import dataclass
from enum import StrEnum

MAX_SIZE = 10 * 1024 * 1024


class DocumentType(StrEnum):
    PDF = "pdf"
    DOCX = "docx"
    TEXT = "txt"
    MARKDOWN = "md"


class TextStatus(StrEnum):
    OK = "ok"
    EMPTY = "empty"  # document lisible, mais sans texte
    UNREADABLE = "unreadable"  # texte non extractible (PDF scanné, fichier abîmé)


CONTENT_TYPES = {
    DocumentType.PDF: "application/pdf",
    DocumentType.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    DocumentType.TEXT: "text/plain; charset=utf-8",
    DocumentType.MARKDOWN: "text/markdown; charset=utf-8",
}


class UnsupportedDocument(ValueError):
    """Type non accepté, contenu qui ne correspond pas à l'extension, ou fichier trop gros."""


@dataclass(frozen=True)
class Extracted:
    type: DocumentType
    status: TextStatus
    text: str


def detect_type(filename: str, data: bytes) -> DocumentType:
    if len(data) > MAX_SIZE:
        raise UnsupportedDocument("fichier trop volumineux (10 Mo au plus)")
    if not data:
        raise UnsupportedDocument("fichier vide")
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension == "pdf":
        if not data.startswith(b"%PDF-"):
            raise UnsupportedDocument("le contenu n'est pas un PDF")
        return DocumentType.PDF
    if extension == "docx":
        if not _is_docx(data):
            raise UnsupportedDocument("le contenu n'est pas un document Word (.docx)")
        return DocumentType.DOCX
    if extension in ("txt", "md"):
        if b"\x00" in data[:4096]:
            raise UnsupportedDocument("le contenu n'est pas du texte")
        return DocumentType.TEXT if extension == "txt" else DocumentType.MARKDOWN
    raise UnsupportedDocument("type non accepté : PDF, DOCX, TXT ou MD")


def _is_docx(data: bytes) -> bool:
    if not data.startswith(b"PK"):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return "word/document.xml" in archive.namelist()
    except zipfile.BadZipFile:
        return False


def _clean(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    # Au plus une ligne vide d'affilée.
    cleaned: list[str] = []
    for line in lines:
        if line or (cleaned and cleaned[-1]):
            cleaned.append(line)
    return "\n".join(cleaned).strip()


def extract(filename: str, data: bytes) -> Extracted:
    doc_type = detect_type(filename, data)
    try:
        match doc_type:
            case DocumentType.PDF:
                text = _pdf_text(data)
            case DocumentType.DOCX:
                text = _docx_text(data)
            case DocumentType.TEXT | DocumentType.MARKDOWN:
                text = _decode(data)
    except Exception:
        return Extracted(doc_type, TextStatus.UNREADABLE, "")
    text = _clean(text)
    if not text:
        # Un PDF sans texte est presque toujours un scan : texte non lisible sans OCR.
        status = TextStatus.UNREADABLE if doc_type is DocumentType.PDF else TextStatus.EMPTY
        return Extracted(doc_type, status, "")
    return Extracted(doc_type, TextStatus.OK, text)


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")
