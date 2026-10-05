"""Documents (extraction, dépôt, suppression) et blocs de profil."""

import io
from collections.abc import AsyncIterator
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from httpx import ASGITransport, AsyncClient
from sqlalchemy import text

from jobbot.api.app import create_app
from jobbot.core.extract import DocumentType, TextStatus, UnsupportedDocument, extract
from jobbot.runtime import Runtime

from .conftest import make_settings


def pdf_with_text(text_line: str) -> bytes:
    """PDF minimal à une page contenant une ligne de texte (police standard)."""
    content = f"BT /F1 12 Tf 72 720 Td ({text_line}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return out.getvalue()


def docx_with(paragraphs: list[str]) -> bytes:
    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Kubernetes"
    table.rows[0].cells[1].text = "CKA en cours"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# --- Extraction -----------------------------------------------------------------------


def test_pdf_text() -> None:
    result = extract("cv.pdf", pdf_with_text("Administrateur systeme Linux"))
    assert (result.type, result.status) == (DocumentType.PDF, TextStatus.OK)
    assert "Administrateur systeme Linux" in result.text


def test_pdf_without_text_is_unreadable() -> None:
    result = extract("scan.pdf", pdf_with_text(""))
    assert (result.status, result.text) == (TextStatus.UNREADABLE, "")


def test_docx_text_with_tables() -> None:
    result = extract("certificat.DOCX", docx_with(["Certificat de travail", "", "", "Bon travail"]))
    assert result.status is TextStatus.OK
    assert result.text == "Certificat de travail\n\nBon travail\nKubernetes | CKA en cours"


def test_text_and_markdown() -> None:
    assert extract("notes.md", b"# Profil\r\nDevOps").text == "# Profil\nDevOps"
    assert extract("vide.txt", b"   \n  ").status is TextStatus.EMPTY
    assert extract("latin.txt", "Genève".encode("latin-1")).text == "Genève"


@pytest.mark.parametrize(
    ("filename", "data", "message"),
    [
        ("cv.exe", b"MZ", "type non accepté"),
        ("faux.pdf", b"pas un pdf", "pas un PDF"),
        ("faux.docx", b"PK\x03\x04pas un zip", "document Word"),
        ("binaire.txt", b"\x00\x01", "pas du texte"),
        ("vide.pdf", b"", "vide"),
        ("gros.txt", b"a" * (10 * 1024 * 1024 + 1), "trop volumineux"),
    ],
)
def test_rejected_documents(filename: str, data: bytes, message: str) -> None:
    with pytest.raises(UnsupportedDocument, match=message):
        extract(filename, data)


# --- API ------------------------------------------------------------------------------


@pytest.fixture
async def api(tmp_path: Path) -> AsyncIterator[AsyncClient]:
    settings = make_settings(storage_path=tmp_path)
    runtime = Runtime.create(settings)
    async with runtime.engine.begin() as conn:
        await conn.execute(text("TRUNCATE documents, profile_chunks CASCADE"))
    app = create_app(settings, runtime)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c
    await runtime.dispose()


async def upload(api: AsyncClient, name: str, data: bytes) -> dict[str, object]:
    response = await api.post("/api/documents", files={"file": (name, data)})
    assert response.status_code == 201, response.text
    body: dict[str, object] = response.json()
    return body


async def test_document_lifecycle(api: AsyncClient, tmp_path: Path) -> None:
    doc = await upload(api, "CV.pdf", pdf_with_text("Ingenieur systeme"))
    assert (doc["doc_type"], doc["text_status"], doc["chunk_count"]) == ("pdf", "ok", 0)
    stored = list((tmp_path / "documents").iterdir())
    assert len(stored) == 1 and stored[0].suffix == ".pdf"

    text_body = (await api.get(f"/api/documents/{doc['id']}/text")).json()
    assert "Ingenieur systeme" in text_body["text"]
    file = await api.get(f"/api/documents/{doc['id']}/file")
    assert file.headers["content-type"] == "application/pdf"
    assert file.content.startswith(b"%PDF-")

    again = await api.post("/api/documents", files={"file": ("copie.pdf", file.content)})
    assert again.status_code == 409
    assert len(list((tmp_path / "documents").iterdir())) == 1

    chunk = await api.post(
        "/api/profile-chunks",
        json={"kind": "experience", "title": "Admin", "content": "Linux", "document_id": doc["id"]},
    )
    assert chunk.status_code == 201
    assert (await api.get("/api/documents")).json()[0]["chunk_count"] == 1

    assert (await api.delete(f"/api/documents/{doc['id']}")).status_code == 204
    assert list((tmp_path / "documents").iterdir()) == []
    # Le bloc reste, sans lien vers le document supprimé.
    [kept] = (await api.get("/api/profile-chunks")).json()
    assert kept["document_id"] is None
    assert (await api.get(f"/api/documents/{doc['id']}/text")).status_code == 404


async def test_rejected_upload(api: AsyncClient) -> None:
    response = await api.post("/api/documents", files={"file": ("x.exe", b"MZ")})
    assert response.status_code == 415
    assert "type non accepté" in response.json()["detail"]


async def test_chunk_crud_and_validation(api: AsyncClient) -> None:
    created = (
        await api.post(
            "/api/profile-chunks",
            json={
                "kind": "competence",
                "title": "  Kubernetes ",
                "content": "CKA en cours, clusters kubeadm",
                "tags": ["infra", "Infra", " k8s "],
            },
        )
    ).json()
    assert (created["title"], created["tags"], created["active"]) == (
        "Kubernetes",
        ["infra", "k8s"],
        True,
    )

    updated = await api.put(
        f"/api/profile-chunks/{created['id']}",
        json={**created, "active": False, "content": "CKA obtenue"},
    )
    assert updated.json()["active"] is False and updated.json()["content"] == "CKA obtenue"

    bad = await api.post(
        "/api/profile-chunks", json={"kind": "autre", "title": "x", "content": "y"}
    )
    assert bad.status_code == 422
    blank = await api.post(
        "/api/profile-chunks", json={"kind": "ton", "title": "   ", "content": "y"}
    )
    assert blank.status_code == 422
    missing_doc = await api.post(
        "/api/profile-chunks",
        json={"kind": "ton", "title": "x", "content": "y", "document_id": 999999},
    )
    assert missing_doc.status_code == 404

    assert (await api.delete(f"/api/profile-chunks/{created['id']}")).status_code == 204
    assert (await api.get("/api/profile-chunks")).json() == []
